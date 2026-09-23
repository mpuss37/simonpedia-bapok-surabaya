#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
generate_uiux_doc.py
Membuat dokumen Word (.docx) berisi dokumentasi UI/UX SIMONPEDIA Bapok Surabaya:
- Judul + catatan awal (disclaimer desain bisa berubah)
- Daftar gambar
- 8 bagian: screenshot + penjelasan fungsional singkat
- Catatan penutup

Jalankan:  python3 docs/generate_uiux_doc.py
Output:    docs/SIMONPEDIA_UIUX_Deskripsi.docx
"""

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

# ----------------------------------------------------------------------------
# KONFIGURASI
# ----------------------------------------------------------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
SS_DIR = os.path.join(BASE_DIR, "screenshots")
OUT_DOCX = os.path.join(BASE_DIR, "SIMONPEDIA_UIUX_Deskripsi.docx")
LOGO = os.path.join(BASE_DIR, "brida-logo.png")

# Warna aksen merah sesuai tema aplikasi
ACCENT = RGBColor(0xC9, 0x37, 0x42)
DARK = RGBColor(0x17, 0x17, 0x17)
GREY = RGBColor(0x66, 0x66, 0x66)

# Lebar gambar konten (cm). A4 lebar 21cm, margin kiri/kanan 2.5cm -> ~16cm.
CONTENT_WIDTH_CM = 15.5
# Tinggi maksimum gambar dalam satu halaman agar tetap muat (cm).
MAX_IMG_HEIGHT_CM = 22.0

# ----------------------------------------------------------------------------
# DATA 8 BAGIAN
# ----------------------------------------------------------------------------
SECTIONS = [
    {
        "no": 1,
        "title": "Dashboard",
        "file": "dashboard-simonpedia-bapok-surabaya.png",
        "subtitle": "Halaman utama / ringkasan kondisi pasar",
        "desc": (
            "Halaman Dashboard adalah tampilan pertama yang dilihat pengguna saat membuka "
            "aplikasi. Halaman ini merangkum kondisi pasar bahan pokok hari itu dalam satu layar, "
            "sehingga pengguna bisa cepat menilai situasi tanpa harus membuka tiap menu."
        ),
        "points": [
            "Bagian hero menampilkan judul \u201cPantau harga bapok Surabaya\u201d beserta tombol pintas menuju Monitoring dan EWS.",
            "Empat kartu ringkasan menampilkan jumlah komoditas dipantau, jumlah harga yang naik, jumlah harga stabil/turun, dan jumlah alert aktif.",
            "Grafik tren harga menampilkan pergerakan harga komoditas terpilih (default: komoditas dengan perubahan terbesar) lengkap dengan garis pembanding HET.",
            "Panel Early Warning menampilkan kondisi risiko terkini, misalnya berapa komoditas yang masuk level Waspada.",
            "Tabel pergerakan komoditas menampilkan harga rata-rata dan persentase perubahan tiap komoditas.",
        ],
    },
    {
        "no": 2,
        "title": "Monitoring Harga",
        "file": "monitoring-simonpedia-bapok-surabaya.png",
        "subtitle": "Pemantauan harga per komoditas",
        "desc": (
            "Halaman Monitoring Harga dipakai untuk memantau harga bahan pokok dari waktu ke waktu. "
            "Di sini pengguna bisa melihat angka ringkas, grafik tren, dan daftar lengkap harga komoditas "
            "beserta status risikonya."
        ),
        "points": [
            "Empat kartu KPI menampilkan rata-rata harga, jumlah komoditas yang naik, turun, dan stabil.",
            "Grafik tren harga dapat difilter berdasarkan periode: 7 Hari, 30 Hari, 3 Bulan, atau 6 Bulan.",
            "Kolom pencarian dan filter kategori membantu menemukan komoditas tertentu dengan cepat.",
            "Tabel harga komoditas menampilkan harga hari ini, persentase perubahan, dan status (Normal, Naik, Turun, Siaga, Waspada, dll).",
            "Setiap baris bisa dibuka untuk melihat detail komoditas.",
        ],
    },
    {
        "no": 3,
        "title": "Daftar Pasar",
        "file": "pasar-simonpedia-bapok-surabaya.png",
        "subtitle": "Daftar pasar amatan",
        "desc": (
            "Halaman Daftar Pasar menampilkan seluruh pasar yang menjadi titik pemantauan harga "
            "bahan pokok di Kota Surabaya. Data pasar ditampilkan sebagai kartu agar mudah dibaca."
        ),
        "points": [
            "Setiap kartu pasar menampilkan nama pasar, kelas (Utama/Madya), kecamatan, wilayah, dan status aktif dipantau.",
            "Filter pencarian memudahkan mencari pasar berdasarkan nama.",
            "Filter kecamatan dan filter kelas membantu mempersempit daftar sesuai kebutuhan.",
            "Jumlah pasar yang ditampilkan ikut berubah mengikuti filter yang aktif.",
        ],
    },
    {
        "no": 4,
        "title": "Peta Pasar",
        "file": "peta-simonpedia-bapok-surabaya.png",
        "subtitle": "Sebaran lokasi pasar di peta",
        "desc": (
            "Halaman Peta Pasar menampilkan sebaran lokasi pasar dalam bentuk peta interaktif. "
            "Pengguna bisa melihat posisi tiap pasar secara geografis sekaligus menelusuri daftarnya."
        ),
        "points": [
            "Peta menampilkan penanda (marker) untuk setiap pasar yang memiliki koordinat lokasi.",
            "Menekan marker akan memunculkan informasi singkat pasar: nama, kecamatan, kelas, dan status.",
            "Dua kartu ringkasan menampilkan total pasar dan jumlah pasar aktif.",
            "Panel daftar di sisi kanan menampilkan pasar terpantau dan ikut menyesuaikan filter pencarian/kecamatan.",
        ],
    },
    {
        "no": 5,
        "title": "Early Warning System (EWS)",
        "file": "ews-simonpedia-bapok-surabaya.png",
        "subtitle": "Deteksi risiko harga",
        "desc": (
            "Halaman Early Warning System (EWS) adalah inti sistem peringatan dini. Sistem memantau "
            "perubahan harga bahan pokok lalu menetapkan tingkat risiko, sehingga potensi kenaikan "
            "harga bisa dideteksi sebelum menjadi masalah besar."
        ),
        "points": [
            "Empat kartu status menampilkan jumlah komoditas pada tiap tingkat risiko: Normal, Siaga, Waspada, dan Kritis.",
            "Banner alert muncul ketika ada komoditas pada level Waspada atau Kritis yang perlu diperhatikan.",
            "Grafik anomali menampilkan perubahan harga komoditas terpilih beserta garis pembanding HET.",
            "Panel daftar alert menampilkan alert terbaru per pasar beserta besaran perubahan harganya.",
            "Ringkasan risiko dan daftar tindakan yang disarankan membantu pengguna menentukan langkah lanjutan.",
        ],
    },
    {
        "no": 6,
        "title": "Rekomendasi Tindakan",
        "file": "rekomendasi-simonpedia-bapok-surabaya.png",
        "subtitle": "Dukungan pengambilan keputusan",
        "desc": (
            "Halaman Rekomendasi Tindakan mengubah hasil analisis data menjadi usulan langkah nyata. "
            "Sistem menyusun rekomendasi berdasarkan perubahan harga, tingkat risiko, dan kondisi pasar."
        ),
        "points": [
            "Kartu ringkasan menampilkan jumlah total rekomendasi, prioritas tinggi, prioritas sedang, dan yang selesai.",
            "Filter Semua / Tinggi / Sedang menyaring daftar rekomendasi sesuai prioritas.",
            "Tiap kartu rekomendasi memuat komoditas, prioritas, harga, perubahan, tindakan yang disarankan, alasan, dan tenggat waktu.",
            "Tersedia tombol aksi seperti \u201cTandai ditangani\u201d dan \u201cDetail analisis\u201d.",
            "Bagian decision engine menjelaskan alur singkat: Monitoring \u2192 Analisis \u2192 Tindakan.",
        ],
    },
    {
        "no": 7,
        "title": "Prediksi Harga",
        "file": "prediksi-simonpedia-bapok-surabaya.png",
        "subtitle": "Perkiraan harga 7 hari ke depan",
        "desc": (
            "Halaman Prediksi Harga memperkirakan harga bahan pokok untuk 7 hari ke depan. "
            "Perhitungan memakai tiga metode sekaligus \u2014 Linear Regression, Moving Average, dan "
            "Exponential Smoothing \u2014 lalu hasilnya diuji akurasinya."
        ),
        "points": [
            "Filter kategori memudahkan memilih kelompok komoditas yang ingin dilihat.",
            "Tabel prediksi menampilkan harga akhir, harga prediksi, persentase perubahan, dan tren (naik/turun/stabil).",
            "Klik salah satu komoditas untuk membuka panel detail prediksi.",
            "Grafik detail menampilkan data historis dan prediksi 7 hari dalam satu tampilan.",
            "Tabel backtest menampilkan nilai MAPE dan RMSE tiap metode, sehingga terlihat metode mana yang paling akurat.",
        ],
    },
    {
        "no": 8,
        "title": "Data & Export",
        "file": "data-simonpedia-bapok-surabaya.png",
        "subtitle": "Unduh data harga",
        "desc": (
            "Halaman Data & Export disediakan untuk pengguna yang ingin mengunduh dan mengolah data "
            "harga bahan pokok sendiri. Pengguna bisa memilih data yang diinginkan dan menentukan "
            "format berkasnya."
        ),
        "points": [
            "Panel filter mengatur komoditas, pasar, dan periode data yang ingin diunduh.",
            "Perkiraan jumlah baris menunjukkan gambaran volume data yang akan diunduh.",
            "Pilihan format berkas: CSV, JSON, dan Excel, masing-masing dengan keterangannya.",
            "Bagian pratinjau menampilkan contoh data (komoditas, kategori, harga rata-rata, perubahan, satuan).",
        ],
    },
]

DISCLAIMER_AWAL = (
    "Dokumen ini memuat dokumentasi tampilan antarmuka (UI/UX) aplikasi SIMONPEDIA Bapok "
    "Surabaya. Perlu dicatat bahwa desain tampilan bersifat sementara dan dapat berubah "
    "sewaktu-waktu, menyesuaikan alur (flow) sistem apabila terjadi perubahan. Namun secara "
    "garis besar, gambaran antarmuka kurang lebih seperti yang ditampilkan pada gambar-gambar "
    "berikut."
)

DISCLAIMER_AKHIR = (
    "Demikian gambaran antarmuka (UI/UX) SIMONPEDIA Bapok Surabaya. Seluruh tampilan pada "
    "dokumen ini masih bersifat sementara dan dapat berubah menyesuaikan alur sistem apabila "
    "terjadi perubahan di kemudian hari. Gambar dan penjelasan di atas dimaksudkan sebagai "
    "gambaran besar, bukan acuan final. Dokumen ini dapat diperbarui seiring perkembangan "
    "aplikasi."
)


# ----------------------------------------------------------------------------
# HELPER
# ----------------------------------------------------------------------------
def set_cell_or_para_spacing(paragraph, before=0, after=6):
    pf = paragraph.paragraph_format
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)


def add_para(doc, text, size=11, bold=False, italic=False, color=None,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=8, font="Calibri"):
    p = doc.add_paragraph()
    p.alignment = align
    set_cell_or_para_spacing(p, before, after)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    if color is not None:
        run.font.color.rgb = color
    return p


def add_heading_custom(doc, text, size=18, color=ACCENT, before=14, after=8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_cell_or_para_spacing(p, before, after)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return p


def add_bullet(doc, text, size=11):
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_cell_or_para_spacing(p, 0, 4)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    return p


def add_hr(doc):
    p = doc.add_paragraph()
    set_cell_or_para_spacing(p, 2, 6)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "C93742")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def img_dimensions(path):
    with Image.open(path) as im:
        return im.size  # (w, h)


def insert_image_fit(doc, path, max_width_cm=CONTENT_WIDTH_CM, max_height_cm=MAX_IMG_HEIGHT_CM):
    """Masukkan gambar dengan rasio terjaga, dibatasi lebar dan tinggi maksimum."""
    w_px, h_px = img_dimensions(path)
    aspect = h_px / w_px  # tinggi / lebar

    width_cm = max_width_cm
    height_cm = width_cm * aspect

    # Kalau terlalu tinggi, kunci tinggi maksimum dan hitung lebar baru.
    if height_cm > max_height_cm:
        height_cm = max_height_cm
        width_cm = height_cm / aspect

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cell_or_para_spacing(p, 4, 4)
    run = p.add_run()
    run.add_picture(path, width=Cm(width_cm), height=Cm(height_cm))
    return width_cm, height_cm


def add_footer():
    pass


# ----------------------------------------------------------------------------
# BANGUN DOKUMEN
# ----------------------------------------------------------------------------
def build():
    doc = Document()

    # --- Gaya default ---
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # Margin halaman A4
    for section in doc.sections:
        section.page_height = Cm(29.7)
        section.page_width = Cm(21.0)
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # ============ HALAMAN JUDUL ============
    for _ in range(3):
        doc.add_paragraph()

    add_para(doc, "DOKUMENTASI ANTARMUKA (UI/UX)",
             size=13, bold=True, color=GREY,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_para(doc, "SIMONPEDIA Bapok Surabaya",
             size=30, bold=True, color=ACCENT,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    add_para(doc, "Sistem Pemantauan Harga Bahan Pokok Kota Surabaya",
             size=13, italic=True, color=DARK,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_para(doc, "Dokumentasi tampilan beserta penjelasan singkat setiap halaman",
             size=11, color=GREY,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=20)

    if os.path.exists(LOGO):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(LOGO, width=Cm(4.5))

    add_para(doc, "Tim EWS", size=12, bold=True, color=DARK,
             align=WD_ALIGN_PARAGRAPH.CENTER, before=20, after=2)
    add_para(doc, "Brida Provinsi Jawa Timur", size=11, color=GREY,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=2)

    doc.add_page_break()

    # ============ CATATAN AWAL ============
    add_heading_custom(doc, "Catatan Awal", size=18, before=0, after=8)
    add_hr(doc)
    add_para(doc, DISCLAIMER_AWAL, size=11)
    add_para(doc, "Dokumen ini disusun untuk memberi gambaran bagaimana tiap halaman "
                  "aplikasi tampil dan apa fungsinya. Setiap gambar diberi nomor dan "
                  "penjelasan singkat di bawahnya.", size=11)

    add_heading_custom(doc, "Daftar Gambar", size=14, before=14, after=6)
    for s in SECTIONS:
        p = doc.add_paragraph()
        set_cell_or_para_spacing(p, 0, 3)
        r1 = p.add_run(f"Gambar {s['no']}. ")
        r1.font.size = Pt(11); r1.font.bold = True; r1.font.color.rgb = ACCENT
        r2 = p.add_run(f"{s['title']} \u2014 {s['subtitle']}")
        r2.font.size = Pt(11)

    doc.add_page_break()

    # ============ 8 BAGIAN ============
    for i, s in enumerate(SECTIONS):
        # Judul bagian
        add_heading_custom(doc, f"Gambar {s['no']}. {s['title']}", size=16,
                           before=0, after=2)
        add_para(doc, s["subtitle"], size=11, italic=True, color=GREY, after=6)
        add_hr(doc)

        img_path = os.path.join(SS_DIR, s["file"])
        if os.path.exists(img_path):
            insert_image_fit(doc, img_path)
            add_para(doc, f"(Gambar {s['no']}: {s['file']})",
                     size=9, italic=True, color=GREY,
                     align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
        else:
            add_para(doc, f"[Gambar tidak ditemukan: {s['file']}]",
                     size=10, italic=True, color=ACCENT)

        # Penjelasan
        add_para(doc, "Penjelasan:", size=11, bold=True, color=DARK,
                 align=WD_ALIGN_PARAGRAPH.LEFT, after=4)
        add_para(doc, s["desc"], size=11)
        for pt in s["points"]:
            add_bullet(doc, pt)

        # Page break antar bagian (kecuali terakhir)
        if i < len(SECTIONS) - 1:
            doc.add_page_break()

    # ============ CATATAN PENUTUP ============
    doc.add_page_break()
    add_heading_custom(doc, "Catatan Penutup", size=18, before=0, after=8)
    add_hr(doc)
    add_para(doc, DISCLAIMER_AKHIR, size=11)

    # Simpan
    doc.save(OUT_DOCX)
    print("OK ->", OUT_DOCX)


if __name__ == "__main__":
    build()
